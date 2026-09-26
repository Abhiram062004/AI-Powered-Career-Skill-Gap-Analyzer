"""
resume_utils.py
---------------
Resume file extraction utilities.

Supported formats:
- PDF
- DOCX
- PNG
- JPG
- JPEG

Features:
- PDF text extraction
- Scanned PDF OCR fallback
- DOCX paragraph extraction
- DOCX table extraction
- Image OCR
- Text cleaning
"""

import io
import os
import re

from PIL import Image, ImageFilter, ImageOps
from PyPDF2 import PdfReader
from docx import Document


# -------------------------------------------------------------------
# Configuration
# -------------------------------------------------------------------

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "png",
    "jpg",
    "jpeg",
}

MAX_TEXT_LENGTH = 100000


# -------------------------------------------------------------------
# File validation
# -------------------------------------------------------------------

def get_file_extension(filename):
    """Return the lowercase extension without the dot."""

    if not filename:
        return ""

    filename = os.path.basename(str(filename))

    if "." not in filename:
        return ""

    return filename.rsplit(".", 1)[1].lower()


def allowed_resume(filename):
    """Check whether the uploaded file has a supported extension."""

    extension = get_file_extension(filename)

    return extension in ALLOWED_EXTENSIONS


def get_file_type(filename):
    """Return a human-readable file type."""

    extension = get_file_extension(filename)

    mapping = {
        "pdf": "PDF",
        "docx": "Microsoft Word",
        "png": "PNG Image",
        "jpg": "JPEG Image",
        "jpeg": "JPEG Image",
    }

    return mapping.get(
        extension,
        "Unknown",
    )


# -------------------------------------------------------------------
# Text cleaning
# -------------------------------------------------------------------

def clean_extracted_text(text):
    """
    Clean text extracted from PDF, DOCX, or OCR.
    """

    if not text:
        return ""

    text = str(text)

    # Normalize line endings.
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove null characters.
    text = text.replace("\x00", " ")

    # Normalize common Unicode spaces.
    text = (
        text.replace("\u00a0", " ")
        .replace("\u2007", " ")
        .replace("\u202f", " ")
    )

    # Normalize unusual dashes.
    text = re.sub(
        r"[\u2010\u2011\u2012\u2013\u2014\u2212]",
        "-",
        text,
    )

    # Remove excessive spaces around lines.
    lines = []

    for line in text.split("\n"):
        line = re.sub(
            r"[ \t]+",
            " ",
            line,
        ).strip()

        if line:
            lines.append(line)

    text = "\n".join(lines)

    # Prevent extremely large extracted text from being stored.
    if len(text) > MAX_TEXT_LENGTH:
        text = text[:MAX_TEXT_LENGTH]

    return text.strip()


# -------------------------------------------------------------------
# PDF extraction
# -------------------------------------------------------------------

def _extract_pdf_text(file_object):
    """
    Extract text from a normal text-based PDF.
    """

    try:
        # Make sure the file starts from the beginning.
        if hasattr(file_object, "seek"):
            file_object.seek(0)

        reader = PdfReader(file_object)

        pages = []

        for page in reader.pages:
            try:
                page_text = page.extract_text() or ""

                if page_text.strip():
                    pages.append(page_text)

            except Exception:
                # Continue extracting remaining pages.
                continue

        return "\n".join(pages)

    except Exception:
        return ""


def _ocr_pdf(file_object):
    """
    OCR scanned PDF pages.

    Requires:
        pdf2image
        pytesseract
        Poppler installed on the system
    """

    try:
        from pdf2image import convert_from_bytes
        import pytesseract
    except ImportError:
        return ""

    try:
        if hasattr(file_object, "seek"):
            file_object.seek(0)

        pdf_bytes = file_object.read()

        if not pdf_bytes:
            return ""

        images = convert_from_bytes(
            pdf_bytes,
            dpi=200,
        )

        extracted_pages = []

        for image in images:
            try:
                processed = preprocess_image(image)

                text = pytesseract.image_to_string(
                    processed
                )

                if text.strip():
                    extracted_pages.append(text)

            except Exception:
                continue

        return "\n".join(extracted_pages)

    except Exception:
        return ""


def extract_pdf_text(file_object):
    """
    Extract PDF text.

    First attempts normal PDF text extraction.
    If little/no text is found, attempts OCR.
    """

    text = _extract_pdf_text(file_object)

    cleaned = clean_extracted_text(text)

    # If a scanned PDF contains no usable text,
    # try OCR.
    if len(cleaned) < 50:
        ocr_text = _ocr_pdf(file_object)

        ocr_cleaned = clean_extracted_text(
            ocr_text
        )

        if len(ocr_cleaned) > len(cleaned):
            cleaned = ocr_cleaned

    return cleaned


# -------------------------------------------------------------------
# DOCX extraction
# -------------------------------------------------------------------

def extract_docx_text(file_object):
    """
    Extract text from DOCX paragraphs and tables.
    """

    try:
        if hasattr(file_object, "seek"):
            file_object.seek(0)

        document = Document(file_object)

        sections = []

        # -----------------------------------------------------------
        # Paragraphs
        # -----------------------------------------------------------

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                sections.append(text)

        # -----------------------------------------------------------
        # Tables
        # -----------------------------------------------------------

        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    sections.append(
                        " | ".join(cells)
                    )

        return clean_extracted_text(
            "\n".join(sections)
        )

    except Exception:
        return ""


# -------------------------------------------------------------------
# Image preprocessing
# -------------------------------------------------------------------

def preprocess_image(image):
    """
    Improve an image before OCR.

    Operations:
    - RGB conversion
    - grayscale
    - contrast enhancement
    - autocontrast
    - sharpening
    - optional upscaling
    """

    if not isinstance(image, Image.Image):
        image = Image.open(image)

    image = image.convert("RGB")

    width, height = image.size

    # Upscale small images for better OCR.
    if width < 1500:
        scale = 2

        image = image.resize(
            (
                width * scale,
                height * scale,
            ),
            Image.Resampling.LANCZOS,
        )

    image = ImageOps.grayscale(image)

    image = ImageOps.autocontrast(
        image
    )

    image = image.filter(
        ImageFilter.SHARPEN
    )

    return image


# -------------------------------------------------------------------
# Image OCR
# -------------------------------------------------------------------

def extract_image_text(file_object):
    """
    Extract text from PNG/JPG/JPEG using Tesseract OCR.
    """

    try:
        import pytesseract
    except ImportError:
        return ""

    try:
        if hasattr(file_object, "seek"):
            file_object.seek(0)

        image = Image.open(file_object)

        processed = preprocess_image(
            image
        )

        text = pytesseract.image_to_string(
            processed
        )

        return clean_extracted_text(
            text
        )

    except Exception:
        return ""


# -------------------------------------------------------------------
# Main extraction function
# -------------------------------------------------------------------

def extract_resume_text(file_object, filename=None):
    """
    Extract text from an uploaded resume.

    Parameters
    ----------
    file_object:
        File-like object.

    filename:
        Original filename.

    Returns
    -------
    str
        Extracted and cleaned text.
    """

    if file_object is None:
        return ""

    if not filename:
        filename = getattr(
            file_object,
            "filename",
            "",
        )

    extension = get_file_extension(
        filename
    )

    if extension == "pdf":
        return extract_pdf_text(
            file_object
        )

    if extension == "docx":
        return extract_docx_text(
            file_object
        )

    if extension in {
        "png",
        "jpg",
        "jpeg",
    }:
        return extract_image_text(
            file_object
        )

    return ""


# -------------------------------------------------------------------
# Alternative API
# -------------------------------------------------------------------

def extract_text(file_object, filename=None):
    """
    Compatibility alias for extract_resume_text().
    """

    return extract_resume_text(
        file_object,
        filename,
    )


def has_resume_text(text, minimum_length=30):
    """
    Check whether extracted resume text is usable.
    """

    if not text:
        return False

    cleaned = clean_extracted_text(
        text
    )

    return len(cleaned) >= minimum_length


# -------------------------------------------------------------------
# Bytes-based helper
# -------------------------------------------------------------------

def extract_text_from_bytes(
    file_bytes,
    filename,
):
    """
    Extract resume text directly from bytes.

    Useful when processing files stored in memory.
    """

    if not file_bytes:
        return ""

    file_object = io.BytesIO(
        file_bytes
    )

    return extract_resume_text(
        file_object,
        filename,
    )


# -------------------------------------------------------------------
# Direct test
# -------------------------------------------------------------------

if __name__ == "__main__":
    print("=" * 60)
    print("Resume Utility Module")
    print("=" * 60)

    print(
        "Supported formats:",
        ", ".join(
            sorted(ALLOWED_EXTENSIONS)
        ),
    )

    print(
        "PDF support      : PyPDF2 + optional OCR"
    )

    print(
        "DOCX support     : python-docx"
    )

    print(
        "Image OCR        : pytesseract"
    )

    print("=" * 60)